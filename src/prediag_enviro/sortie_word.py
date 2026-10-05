"""Construction du document Word — un rapport, pas un export de tableaux.

La première version produisait des tableaux corrects dans un document vide.
Mis à côté du prédiagnostic de référence, l'écart ne portait pas sur les
données mais sur tout ce qui en fait un rapport : des sections, des phrases de
constat, des légendes numérotées, et dans les tableaux une **ligne fusionnée
par type** qui groupe les entrées au lieu d'une colonne « Type » répétée.

**Le document part d'un modèle, pas d'une page blanche.** `modele/`
porte un `.docx` tiré du prédiagnostic de référence, vidé de son contenu : il
apporte l'en-tête au logo, le pied de page avec sa pagination, les marges, et
surtout les styles maison — « Titre partie », « Titre colonne », « Corps de
texte - Unite ». Les recréer à l'approchant, comme le faisait la première
version, donnait un document qui ne ressemblait à rien de ce que l'équipe
produit. Un modèle déposé par l'utilisatrice prend le pas sur celui-ci.

`Rapport` tient les compteurs et les gestes du document. Les compteurs
comptent : une légende « Tableau 8 » puis « Tableau 9 » se tient toute seule,
alors que « Tableau : » répété douze fois oblige à tout renuméroter à la main —
exactement le geste que l'outil est censé supprimer. Les numéros sont en outre
posés dans des **champs `SEQ`**, ce qui permet à Word de les recalculer et
surtout d'en tirer les sommaires des tableaux et des cartes.

Les conventions de mise en forme viennent de l'onglet « Utilisation » des
classeurs, qui décrit le geste manuel qu'on remplace :

    Copier le tableau (sans prendre les cellules en violet) · Coller dans Word
    et appliquer les différents styles du tableau (texte tableau centré ou non
    et titre colonne tableau), réajuster les colonnes et mettre les noms latins
    en italique

===================  =========================================================
titre du document    style « Title »
section              style « Titre partie »
sous-section         style « Style1 »
en-tête de colonne   style « Titre colonne »
corps de cellule     style « Corps de texte - Unite »
nom scientifique     en italique
dates et codes       centrés
tableau              style « Table Grid »
légende              style « Caption »
===================  =========================================================
"""
from __future__ import annotations

import re
from collections.abc import Callable
from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK,
                            WD_COLOR_INDEX)
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import chemins

STYLE_TITRE = "Title"
STYLE_PARTIE = "Titre partie"
STYLE_SOUS_PARTIE = "Style1"
#: Troisième niveau. Le prédiagnostic externe en a besoin : son état
#: initial descend jusqu'à « ZNIEFF de type I et II et ZICO ».
STYLE_SOUS_SOUS_PARTIE = "Style2"
STYLE_ENTETE = "Titre colonne"
STYLE_CORPS = "Corps de texte - Unite"
STYLE_LEGENDE = "Caption"
STYLE_TABLEAU = "Table Grid"
STYLE_PUCE = "List Paragraph"
STYLE_NORMAL = "Normal"
STYLE_SOMMAIRE = "TOC Heading"

#: Couleur du texte des consignes laissées à l'expert, posée par-dessus un
#: surlignage jaune. La première version se contentait de la couleur, au motif
#: qu'un surlignage oublié se voit mal à l'écran ; à l'usage, c'est l'inverse
#: qui gêne — une couleur de caractère ne saute pas aux yeux quand on feuillette
#: quarante pages pour repérer ce qui reste à faire.
COULEUR_CONSIGNE = RGBColor(0x1F, 0x4E, 0x79)

#: Gris des mentions secondaires de la page de garde.
COULEUR_DISCRETE = RGBColor(0x59, 0x59, 0x59)

# ════════════════════════════════════════════════ palette des tableaux ══
#
# Relevée dans le prédiagnostic de référence plutôt que choisie : les tableaux
# n'ont **pas de traits** — leurs bordures sont blanches — et c'est l'ombrage
# des cellules qui les structure. Reproduire « Table Grid » et ses filets noirs
# donnait des tableaux que personne de l'équipe n'aurait reconnus.

#: En-tête et première colonne : le vert pâle qui identifie la ligne.
VERT_ENTETE = "E2EFD9"
#: Corps des tableaux.
CREME_CORPS = "EEECE1"

#: Fond par catégorie de liste rouge. Mesuré sur 243 cellules du document de
#: référence : LC vert (127 cellules), NT jaune pâle (45), VU jaune (24),
#: EN orange (7), CR rouge (5), NA gris (32), DD gris clair (3).
#:
#: `RE`, `EX` et `CR*` n'y figurent pas — aucune espèce du document n'est dans
#: ce cas. Ils prennent le rouge de `CR`, dont ils sont les degrés supérieurs :
#: c'est une extrapolation, mais les laisser en crème sous-signalerait la
#: catégorie la plus grave de l'échelle.
COULEURS_STATUT = {
    "LC": "78B74A",
    "NT": "FBF2CA",
    "VU": "FFED00",
    "EN": "FBBF00",
    "CR": "D3001B", "CR*": "D3001B", "RE": "D3001B", "EX": "D3001B",
    "DD": "D3D4D5", "NE": "D3D4D5",
    "NA": "929395", "NAa": "929395", "NAb": "929395",
    "NAc": "929395", "NAd": "929395",
}

#: Fonds sur lesquels le texte passe en blanc. Le document de référence ne le
#: fait que sur le gris — le rouge `CR` y garde un texte noir.
FONDS_SOMBRES = {"929395"}

#: Marge intérieure des cellules, en vingtièmes de point. Serrée : les tableaux
#: d'espèces ont jusqu'à dix colonnes sur une page A4.
MARGE_CELLULE = 70


def _ombrer(cellule, couleur: str) -> None:
    from docx.oxml import OxmlElement

    proprietes = cellule._tc.get_or_add_tcPr()
    for ancien in proprietes.findall(qn("w:shd")):
        proprietes.remove(ancien)
    ombre = OxmlElement("w:shd")
    ombre.set(qn("w:val"), "clear")
    ombre.set(qn("w:color"), "auto")
    ombre.set(qn("w:fill"), couleur)
    proprietes.append(ombre)


def _aligner_verticalement(cellule, position: str) -> None:
    from docx.oxml import OxmlElement

    proprietes = cellule._tc.get_or_add_tcPr()
    for ancien in proprietes.findall(qn("w:vAlign")):
        proprietes.remove(ancien)
    noeud = OxmlElement("w:vAlign")
    noeud.set(qn("w:val"), position)
    proprietes.append(noeud)


def _repeter_en_tete(rangee) -> None:
    """Rejoue la ligne d'en-tête en haut de chaque page.

    Un tableau d'avifaune fait couramment quarante lignes et déborde sur deux
    pages ; sans cela, la seconde page porte dix colonnes de codes sans titre.
    """
    from docx.oxml import OxmlElement

    proprietes = rangee._tr.get_or_add_trPr()
    proprietes.append(OxmlElement("w:tblHeader"))


def _bordures_blanches(tableau) -> None:
    from docx.oxml import OxmlElement

    proprietes = tableau._tbl.tblPr
    for ancien in proprietes.findall(qn("w:tblBorders")):
        proprietes.remove(ancien)
    bordures = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        trait = OxmlElement(f"w:{cote}")
        trait.set(qn("w:val"), "single")
        trait.set(qn("w:sz"), "4")
        trait.set(qn("w:space"), "0")
        trait.set(qn("w:color"), "FFFFFF")
        bordures.append(trait)
    proprietes.append(bordures)


def _largeur_pleine(tableau, colonnes) -> None:
    """Table à 100 % de la justification, colonnes en pourcentage.

    Le document de référence exprime ses largeurs en pourcentage et non en
    centimètres : la même table reste correcte en portrait comme en paysage.
    """
    from docx.oxml import OxmlElement

    proprietes = tableau._tbl.tblPr
    for ancien in proprietes.findall(qn("w:tblW")):
        proprietes.remove(ancien)
    largeur = OxmlElement("w:tblW")
    largeur.set(qn("w:w"), "5000")
    largeur.set(qn("w:type"), "pct")
    proprietes.insert(0, largeur)

    marges = OxmlElement("w:tblCellMar")
    for cote in ("left", "right"):
        noeud = OxmlElement(f"w:{cote}")
        noeud.set(qn("w:w"), str(MARGE_CELLULE))
        noeud.set(qn("w:type"), "dxa")
        marges.append(noeud)
    proprietes.append(marges)

    total = sum(c.largeur_cm or 0 for c in colonnes)
    if total <= 0:
        return
    for rangee in tableau.rows:
        if len(rangee.cells) != len(colonnes):
            continue   # ligne fusionnée : la largeur ne s'y applique pas
        for cellule, colonne in zip(rangee.cells, colonnes):
            part = int(round((colonne.largeur_cm or 0) / total * 5000))
            proprietes_cellule = cellule._tc.get_or_add_tcPr()
            for ancien in proprietes_cellule.findall(qn("w:tcW")):
                proprietes_cellule.remove(ancien)
            noeud = OxmlElement("w:tcW")
            noeud.set(qn("w:w"), str(part))
            noeud.set(qn("w:type"), "pct")
            proprietes_cellule.insert(0, noeud)


def modele_par_defaut() -> Path:
    """Le modèle livré avec l'outil, à côté du code."""
    return chemins.racine_application() / "modele" / "prediag_modele.docx"


#: Colonnes dont le contenu se centre (codes courts, dates, statuts).
_CENTREES = {
    "dernière observation", "nidification", "reproduction", "distance à la zip",
    "identifiant", "aire(s) d'étude concernée(s)", "note /10", "cd_nom",
    "protection nationale", "protection régionale", "directive habitats",
    "liste rouge europe", "liste rouge nationale", "liste rouge régionale",
    "liste rouge france", "liste rouge mondiale", "nicheurs", "hivernants",
    "de passage", "annexe i de la directive oiseaux",
    "annexe ii de la directive habitats", "plan national d'actions",
}


@dataclass
class Colonne:
    titre: str
    cle: str                 # clé dans le dictionnaire de ligne
    italique: bool = False   # noms scientifiques
    centree: bool | None = None
    largeur_cm: float | None = None

    def alignee(self) -> bool:
        if self.centree is not None:
            return self.centree
        return self.titre.strip().lower() in _CENTREES


#: « Tableau 3 : ... », « Carte : ... » — le numéro est facultatif, le document
#: de référence mêle des légendes numérotées et des légendes qui ne le sont pas.
_ETIQUETTE = re.compile(r"^\s*(Tableau|Carte|Figure)\s*\d*\s*[:\u00a0:]\s*(.*)$")


def _nom_style(paragraphe) -> str:
    """Identifiant de style d'un `w:p`, ou chaîne vide."""
    proprietes = paragraphe.find(qn("w:pPr"))
    if proprietes is None:
        return ""
    noeud = proprietes.find(qn("w:pStyle"))
    return noeud.get(qn("w:val")) if noeud is not None else ""


def _vider(document: Document) -> Document:
    """Ne garde du modèle que la dernière `sectPr` — sa mise en page."""
    corps = document.element.body
    for enfant in list(corps):
        if not enfant.tag.endswith("}sectPr"):
            corps.remove(enfant)
    return document


# ────────────────────────────────────────────────────────────────── champs ──

def _element(balise: str, **attributs):
    from docx.oxml import OxmlElement

    noeud = OxmlElement(balise)
    for nom, valeur in attributs.items():
        noeud.set(qn(nom.replace("_", ":")), valeur)
    return noeud


def _champ(paragraphe, instruction: str, resultat: str = "",
           recalculer: bool = False) -> None:
    """Insère un champ Word, avec son résultat déjà calculé.

    Le résultat compte autant que l'instruction. Un champ sans résultat
    s'affiche vide tant que personne n'a pressé F9 : un document relu dans un
    convertisseur, un aperçu ou un export PDF automatique montrerait des
    légendes « Tableau  : » sans numéro. On écrit donc le numéro qu'on a
    calculé, et Word le recalcule s'il le juge utile.
    """
    debut = _element("w:fldChar", w_fldCharType="begin")
    if recalculer:
        debut.set(qn("w:dirty"), "true")
    paragraphe._p.append(_passage(debut))

    texte = _element("w:instrText")
    texte.set(qn("xml:space"), "preserve")
    texte.text = instruction
    paragraphe._p.append(_passage(texte))

    paragraphe._p.append(_passage(_element("w:fldChar", w_fldCharType="separate")))
    if resultat:
        noeud = _element("w:t")
        noeud.set(qn("xml:space"), "preserve")
        noeud.text = resultat
        paragraphe._p.append(_passage(noeud))
    paragraphe._p.append(_passage(_element("w:fldChar", w_fldCharType="end")))


def _passage(enfant):
    """Enveloppe un nœud dans un `w:r`, seul conteneur admis dans un `w:p`."""
    from docx.oxml import OxmlElement

    run = OxmlElement("w:r")
    run.append(enfant)
    return run


class Rapport:
    """Un document en construction.

    Tous les gestes passent par ici pour que la numérotation des tableaux et
    des cartes reste cohérente : elle ne peut pas l'être si les appelants
    écrivent leurs légendes eux-mêmes.
    """

    def __init__(self, modele: Path | None = None):
        choisi = Path(modele) if modele else modele_par_defaut()
        self.modele_utilise = choisi if choisi.exists() else None
        if self.modele_utilise is not None:
            self.document = _vider(Document(str(self.modele_utilise)))
        else:
            self.document = Document()
        self.styles_recrees = self._assurer_styles()
        self._n_tableau = 0
        self._n_carte = 0

    # ------------------------------------------------------------------ styles

    def _assurer_styles(self) -> list[str]:
        """Crée les styles manquants. Renvoie ceux qui ont dû l'être."""
        from docx.enum.style import WD_STYLE_TYPE

        crees: list[str] = []
        existants = {s.name for s in self.document.styles}
        # (nom, taille, gras, espace avant, espace après)
        voulus = (
            (STYLE_PARTIE, 14, True, 12, 6),
            (STYLE_SOUS_PARTIE, 11, True, 8, 4),
            (STYLE_SOUS_SOUS_PARTIE, 10, True, 6, 3),
            (STYLE_ENTETE, 9, True, 1, 1),
            (STYLE_CORPS, 9, False, 1, 1),
        )
        for nom, taille, gras, avant, apres in voulus:
            if nom in existants:
                continue
            style = self.document.styles.add_style(nom, WD_STYLE_TYPE.PARAGRAPH)
            style.font.size = Pt(taille)
            style.font.bold = gras
            style.paragraph_format.space_before = Pt(avant)
            style.paragraph_format.space_after = Pt(apres)
            crees.append(nom)
        return crees

    def _style(self, paragraphe, nom: str):
        try:
            paragraphe.style = nom
        except KeyError:
            pass
        return paragraphe

    # ------------------------------------------------------------------ titres

    def titre(self, texte: str) -> None:
        self._style(self.document.add_paragraph(texte), STYLE_TITRE)

    def partie(self, texte: str) -> None:
        self._style(self.document.add_paragraph(texte), STYLE_PARTIE)

    def sous_partie(self, texte: str) -> None:
        self._style(self.document.add_paragraph(texte), STYLE_SOUS_PARTIE)

    def sous_sous_partie(self, texte: str) -> None:
        self._style(self.document.add_paragraph(texte), STYLE_SOUS_SOUS_PARTIE)

    # ------------------------------------------------------------------ textes

    def para(self, texte: str = "") -> None:
        self._style(self.document.add_paragraph(texte), STYLE_NORMAL)

    def puces(self, elements: list[str]) -> None:
        for element in elements:
            paragraphe = self.document.add_paragraph(element)
            self._style(paragraphe, STYLE_PUCE)
            paragraphe.paragraph_format.left_indent = Cm(1)

    def note(self, texte: str) -> None:
        """« Note : … » — le registre qu'emploie le document de référence."""
        paragraphe = self.document.add_paragraph()
        self._style(paragraphe, STYLE_NORMAL)
        passage = paragraphe.add_run("Note : ")
        passage.italic = True
        paragraphe.add_run(texte).italic = True

    def a_completer(self, consigne: str) -> None:
        """Un trou à combler par l'expert, et ce qu'on attend précisément.

        Écrit en couleur et préfixé : le document doit pouvoir être parcouru
        une fois pour repérer tout ce qui reste à faire. Une consigne vague
        (« à compléter ») ferait relire le document entier pour savoir quoi
        écrire, donc on dit ce qu'on attend.
        """
        paragraphe = self.document.add_paragraph()
        self._style(paragraphe, STYLE_NORMAL)
        passage = paragraphe.add_run("[À compléter — " + consigne + "]")
        # Surlignage jaune : une couleur de caractère seule ne se repère pas
        # au feuilletage d'un document de quarante pages, et c'est justement
        # ce qu'on demande au relecteur de faire.
        passage.font.highlight_color = WD_COLOR_INDEX.YELLOW
        passage.font.color.rgb = COULEUR_CONSIGNE
        passage.italic = True
        passage.bold = True

    def saut_de_page(self) -> None:
        self.document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ──────────────────────────────────────────── page de garde et sommaires ──

    def page_de_garde(self, titre: str, sous_titre: str = "",
                      lignes: list[tuple[str, str]] | None = None,
                      mention: str = "") -> None:
        """La première page : ce qu'est le document et sur quoi il porte.

        L'en-tête au logo et le pied de page paginé du modèle s'y appliquent
        comme sur le reste du document — c'est le parti du prédiagnostic de
        référence, qui ne distingue pas sa première page.
        """
        for _ in range(3):
            self.document.add_paragraph()

        bloc = self.document.add_paragraph()
        bloc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        passage = bloc.add_run(titre)
        passage.bold = True
        passage.font.size = Pt(28)

        if sous_titre:
            bas = self.document.add_paragraph()
            bas.alignment = WD_ALIGN_PARAGRAPH.CENTER
            passage = bas.add_run(sous_titre)
            passage.font.size = Pt(16)

        self.document.add_paragraph()
        for intitule, valeur in (lignes or []):
            ligne = self.document.add_paragraph()
            ligne.alignment = WD_ALIGN_PARAGRAPH.CENTER
            gras = ligne.add_run(f"{intitule} : ")
            gras.bold = True
            gras.font.size = Pt(10)
            ligne.add_run(valeur).font.size = Pt(10)

        if mention:
            for _ in range(2):
                self.document.add_paragraph()
            pied = self.document.add_paragraph()
            pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
            passage = pied.add_run(mention)
            passage.italic = True
            passage.font.size = Pt(9)
            passage.font.color.rgb = COULEUR_DISCRETE

        self.saut_de_page()

    def _titre_sommaire(self, texte: str) -> None:
        paragraphe = self.document.add_paragraph(texte)
        if not self._style(paragraphe, STYLE_SOMMAIRE).style.name == STYLE_SOMMAIRE:
            self._style(paragraphe, STYLE_PARTIE)

    def sommaire(self, avec_tableaux: bool = True,
                 avec_cartes: bool = True) -> None:
        """Sommaire des parties, puis des tableaux et des cartes.

        Les trois sont des champs : Word les remplit à l'ouverture, parce que
        seuls les numéros de page le renseignent et qu'ils dépendent de la mise
        en page finale. Le texte de remplacement dit quoi faire si le document
        est ouvert dans un éditeur qui ne calcule pas les champs.

        Les sommaires des tableaux et des cartes reposent sur les champs `SEQ`
        posés dans les légendes : sans eux, Word ne saurait pas distinguer un
        tableau d'une carte.
        """
        consigne = "Sommaire à mettre à jour : Ctrl+A puis F9."

        self._titre_sommaire("Sommaire")
        paragraphe = self.document.add_paragraph()
        _champ(paragraphe, r' TOC \o "1-2" \h \z \u ', consigne, recalculer=True)

        if avec_tableaux:
            self._titre_sommaire("Sommaire des tableaux")
            paragraphe = self.document.add_paragraph()
            _champ(paragraphe, r' TOC \h \z \c "Tableau" ', consigne,
                   recalculer=True)

        if avec_cartes:
            self._titre_sommaire("Sommaire des cartes")
            paragraphe = self.document.add_paragraph()
            _champ(paragraphe, r' TOC \h \z \c "Carte" ', consigne,
                   recalculer=True)

        self.saut_de_page()

    def _legende(self, etiquette: str, numero: int, texte: str,
                 centree: bool = False, paragraphe=None):
        """« Tableau 3 : … », dont le numéro est un champ `SEQ`.

        `paragraphe` permet de réécrire une légende existante — celles que
        porte le document de cadre, dont la numérotation doit s'enchaîner avec
        celle du rapport.
        """
        if paragraphe is None:
            paragraphe = self.document.add_paragraph()
        else:
            # On vide le paragraphe de ses passages en gardant ses propriétés,
            # sinon l'ancienne légende se cumulerait avec la nouvelle.
            for enfant in list(paragraphe._p):
                if not enfant.tag.endswith("}pPr"):
                    paragraphe._p.remove(enfant)
        self._style(paragraphe, STYLE_LEGENDE)
        if centree:
            paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraphe.add_run(f"{etiquette} ")
        _champ(paragraphe, f" SEQ {etiquette} \\* ARABIC ", str(numero))
        paragraphe.add_run(f" : {texte}")
        return paragraphe

    # ────────────────────────────────────────────────── insertion d'un cadre

    def inserer(self, chemin: Path,
                reperes: dict[str, "Callable[[], None]"] | None = None) -> int:
        """Insère le corps d'un autre document dans celui-ci.

        Sert au **cadre général** : quatre-vingts paragraphes de contenu
        réglementaire — définition des ZNIEFF, article L. 411-1, outils de
        bioévaluation — identiques d'un prédiagnostic à l'autre. Ils vivent
        dans un `.docx` que l'on édite dans Word, et non dans le code : la
        personne qui les fera évoluer travaille avec la réglementation, pas
        avec Python.

        `reperes` associe un intertitre du cadre à une fonction appelée juste
        après lui. C'est ainsi que « Définition des aires d'étude » reçoit le
        tableau des aires réellement retenues pour le projet, sans qu'il faille
        découper le cadre en plusieurs fichiers.

        Les légendes du cadre sont **renumérotées** : elles doivent s'enchaîner
        avec celles du rapport, sans quoi les sommaires des tableaux et des
        cartes compteraient deux fois à partir de un.

        Renvoie le nombre d'éléments insérés.
        """
        chemin = Path(chemin)
        if not chemin.exists():
            return 0
        source = Document(str(chemin))
        styles = {s.style_id: s.name for s in source.styles}
        corps = self.document.element.body
        fin = corps.find(qn("w:sectPr"))
        reperes = reperes or {}
        inseres = 0

        for element in source.element.body.iterchildren():
            balise = element.tag.split("}")[-1]
            if balise not in ("p", "tbl"):
                continue
            copie = deepcopy(element)
            if fin is not None:
                corps.insert(list(corps).index(fin), copie)
            else:
                corps.append(copie)
            inseres += 1
            if balise != "p":
                continue

            paragraphe = Paragraph(copie, self.document)
            nom_style = styles.get(_nom_style(copie), "")
            texte = paragraphe.text.strip()

            if nom_style in ("Lgende", "Caption"):
                trouve = _ETIQUETTE.match(texte)
                if trouve:
                    etiquette = trouve.group(1)
                    if etiquette == "Carte":
                        self._n_carte += 1
                        numero = self._n_carte
                    else:
                        self._n_tableau += 1
                        numero = self._n_tableau
                    self._legende(etiquette, numero, trouve.group(2).strip(),
                                  paragraphe=paragraphe)
            elif texte in reperes:
                # Le repère est déjà posé ; ce qui suit s'ajoute à la fin du
                # corps, donc avant la `sectPr`, c'est-à-dire juste après lui.
                reperes[texte]()

        return inseres

    # --------------------------------------------------------------- tableaux

    def _cellule(self, cellule, texte: str, style: str, italique: bool,
                 centree: bool, gras: bool = False,
                 fond: str | None = None) -> None:
        paragraphe = cellule.paragraphs[0]
        self._style(paragraphe, style)
        if centree:
            paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
        valeur = str(texte if texte is not None else "")
        passage = paragraphe.add_run(valeur)
        passage.italic = italique
        if gras:
            passage.bold = True

        # Le fond dit la catégorie avant qu'on ait lu le code : une colonne de
        # listes rouges se parcourt à la couleur.
        statut = COULEURS_STATUT.get(valeur.strip())
        couleur = statut or fond
        if couleur:
            _ombrer(cellule, couleur)
        if couleur in FONDS_SOMBRES:
            passage.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    def tableau(self, colonnes: list[Colonne], lignes: list[dict],
                legende: str = "",
                groupes: list[tuple[str, list[dict]]] | None = None,
                legende_bas: str = "") -> int:
        """Ajoute un tableau. Renvoie son numéro, pour y renvoyer dans le texte.

        `groupes` reproduit la forme du document de référence : chaque élément
        est un couple (sous-en-tête, lignes), et le sous-en-tête occupe une
        ligne fusionnée sur toute la largeur. C'est ce qui permet d'écrire
        « ZNIEFF de type I » une fois au-dessus de ses onze entrées plutôt
        qu'une colonne « Type » qui répète la même valeur onze fois.

        `legende_bas` porte la ligne « Légende : … » que le document place sous
        les tableaux dont les colonnes emploient des sigles.
        """
        self._n_tableau += 1
        numero = self._n_tableau

        if legende:
            self._legende("Tableau", numero, legende)

        tableau = self.document.add_table(rows=1, cols=len(colonnes))
        try:
            tableau.style = STYLE_TABLEAU
        except KeyError:
            pass
        tableau.alignment = WD_TABLE_ALIGNMENT.CENTER

        entete = tableau.rows[0]
        _repeter_en_tete(entete)
        for cellule, colonne in zip(entete.cells, colonnes):
            self._cellule(cellule, colonne.titre, STYLE_ENTETE, False, True,
                          fond=VERT_ENTETE)
            _aligner_verticalement(cellule, "center")

        def _ajouter(ligne: dict) -> None:
            cellules = tableau.add_row().cells
            for rang, (cellule, colonne) in enumerate(zip(cellules, colonnes)):
                # La première colonne garde le vert de l'en-tête : c'est elle
                # qui identifie la ligne, et l'œil la suit de bout en bout.
                self._cellule(cellule, ligne.get(colonne.cle, ""), STYLE_CORPS,
                              colonne.italique, colonne.alignee(),
                              fond=VERT_ENTETE if rang == 0 else CREME_CORPS)

        if groupes:
            for sous_entete, sous_lignes in groupes:
                if not sous_lignes:
                    continue
                rangee = tableau.add_row()
                fusionnee = rangee.cells[0].merge(rangee.cells[-1])
                # La cellule fusionnée hérite des paragraphes de toutes les
                # cellules d'origine ; on ne garde que le premier, sinon le
                # sous-en-tête s'affiche suivi d'autant de lignes vides qu'il y
                # avait de colonnes.
                for paragraphe in list(fusionnee.paragraphs[1:]):
                    paragraphe._element.getparent().remove(paragraphe._element)
                self._cellule(fusionnee, sous_entete, STYLE_ENTETE, False, True,
                              gras=True, fond=VERT_ENTETE)
                _aligner_verticalement(fusionnee, "center")
                for ligne in sous_lignes:
                    _ajouter(ligne)
        else:
            for ligne in lignes:
                _ajouter(ligne)

        _bordures_blanches(tableau)
        _largeur_pleine(tableau, colonnes)

        if legende_bas:
            self._style(self.document.add_paragraph(legende_bas), STYLE_LEGENDE)
        return numero

    # ----------------------------------------------------------------- cartes

    def _section(self, paysage: bool) -> None:
        """Ouvre une nouvelle section, portrait ou paysage.

        Les cartes sont produites en A4 paysage. Les poser en portrait obligeait
        à les réduire à 16 cm de large, soit la moitié de leur définition utile :
        les toponymes du fond de plan devenaient illisibles. Le prédiagnostic de
        référence fait de même — il alterne vingt et une sections pour que
        chaque carte occupe une page couchée.

        L'en-tête et le pied de page suivent : une nouvelle section hérite des
        références de la précédente tant qu'on ne les délie pas.
        """
        section = self.document.add_section(WD_SECTION.NEW_PAGE)
        largeur, hauteur = section.page_width, section.page_height
        voulu_paysage = largeur > hauteur
        if voulu_paysage != paysage:
            section.orientation = (WD_ORIENT.LANDSCAPE if paysage
                                   else WD_ORIENT.PORTRAIT)
            section.page_width, section.page_height = hauteur, largeur

    def carte(self, chemin: Path, legende: str = "",
              largeur_cm: float | None = None, paysage: bool = True) -> int:
        """Insère une carte sur sa propre page, avec sa légende numérotée."""
        self._n_carte += 1
        if paysage:
            self._section(paysage=True)
        largeur = largeur_cm or (24.0 if paysage else 16.0)

        paragraphe = self.document.add_paragraph()
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraphe.add_run().add_picture(str(chemin), width=Cm(largeur))
        if legende:
            self._legende("Carte", self._n_carte, legende, centree=True)
        if paysage:
            self._section(paysage=False)
        return self._n_carte

    # ------------------------------------------------------------ pied, sortie

    def sources(self, lignes: list[str]) -> None:
        """Sources et dates de consultation, en fin de document.

        La Licence Ouverte 2.0 sous laquelle l'INPN diffuse ses référentiels
        impose de citer la source **et sa date** : ce bloc n'est pas une
        politesse, c'est la condition d'usage de la donnée.
        """
        if not lignes:
            return
        # Dédoublonnage : les référentiels de zonages sont cités à la fois par
        # le contexte et par la préparation des zonages, et les appelants n'ont
        # pas à le savoir. L'ordre de première apparition est conservé.
        vues: list[str] = []
        for ligne in lignes:
            if ligne and ligne not in vues:
                vues.append(ligne)
        self.partie("Sources")
        for ligne in vues:
            paragraphe = self.document.add_paragraph(ligne)
            self._style(paragraphe, STYLE_NORMAL)
            for passage in paragraphe.runs:
                passage.font.size = Pt(8)

    def _recalculer_champs(self) -> None:
        """Demande à Word de mettre les champs à jour à l'ouverture.

        Sans cela, le sommaire reste sur la consigne de remplacement jusqu'à
        ce que quelqu'un pense à presser F9 — et personne n'y pense.
        """
        parametres = self.document.settings.element
        if parametres.find(qn("w:updateFields")) is None:
            noeud = _element("w:updateFields", w_val="true")
            parametres.append(noeud)

    def enregistrer(self, chemin: Path) -> Path:
        self._recalculer_champs()
        chemin = Path(chemin)
        chemin.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(chemin))
        return chemin


# ------------------------------------------------------------------- colonnes

def colonnes_zonages() -> list[Colonne]:
    """Colonnes des tableaux de zonages, dans l'ordre du document de référence.

    Pas de colonne « Type » : le type devient la ligne fusionnée qui groupe les
    entrées, comme dans le document de référence.
    """
    return [
        Colonne("Nom", "nom", largeur_cm=4.2),
        Colonne("Distance à la ZIP", "distance", centree=True, largeur_cm=2.2),
        Colonne("Identifiant", "identifiant", centree=True, largeur_cm=2.2),
        Colonne("Intérêt", "interet", largeur_cm=5.4),
        Colonne("Aire(s) d'étude concernée(s)", "aires", centree=True,
                largeur_cm=2.0),
    ]


def colonnes_risque() -> list[Colonne]:
    """Colonnes du tableau de note de risque, en conclusion."""
    return [
        Colonne("Groupe", "groupe", largeur_cm=3.2),
        Colonne("Note /10", "note", centree=True, largeur_cm=1.8),
        Colonne("Commentaires", "commentaires", largeur_cm=11.0),
    ]
