"""Lecture des classeurs B-Statuts, en LECTURE SEULE.

Les six classeurs n'ont pas la même disposition : CD_NOM est en colonne 4 pour
les oiseaux, 5 pour les mammifères et les araignées, 6 pour les insectes. Les
en-têtes de l'onglet `Statuts` sont par ailleurs fusionnés (« Liste rouge
France » couvre trois colonnes : nicheurs, hivernants, de passage) et ne disent
donc pas à quelle période chaque colonne correspond.

La solution ne vient pas d'une heuristique mais du classeur lui-même :
**l'onglet « Saisie étude » est un dictionnaire de colonnes** que Marie tient à
jour. Trois lignes le composent :

    ligne 5 : l'index de la colonne dans l'onglet `Statuts`
    ligne 6 : le territoire        (Europe, France, Normandie…)
    ligne 7 : le champ             (LRE, PN, LRF Nicheurs, LRR Hivernants…)

On lit ce dictionnaire plutôt que de deviner. S'il manque, on retombe sur les
en-têtes de `Statuts` et leurs plages fusionnées, en signalant les colonnes dont
la période reste indéterminée.

Rien n'est jamais écrit dans ces fichiers : voir `veille.py` pour les raisons.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# Champ du classeur -> type de statut BDC-Statuts.
CHAMP_VERS_BDC = {
    "LRE": "LRE",
    "LRM": "LRM",
    "LRF": "LRN",
    "LRN": "LRN",
    "LRR": "LRR",
    "DZ": "ZDET",
    "ZDET": "ZDET",
    "PN": "PN",
    "PR": "PR",
    "PD": "PD",
    "ANN. I DO": "DO",
    "ANN. II DH": "DH",
    "ANN. II ET IV DH": "DH",
    "DH": "DH",
    "DO": "DO",
    "PNA": "PNA",
}

# Territoires nommés différemment chez Marie et dans BDC-Statuts.
ALIAS_TERRITOIRE = {
    "centre val de loire": "Centre",
    "ile de france": "Ile-de-France",
    "pays de la loire": "Pays-de-la-Loire",
    "provence alpes cote d azur": "Provence-Alpes-Côte-d'Azur",
    "nord pas de calais": "Nord-Pas-de-Calais",
    "languedoc roussillon": "Languedoc-Roussillon",
    "midi pyrenees": "Midi-Pyrénées",
    "poitou charentes": "Poitou-Charentes",
    "champagne ardenne": "Champagne-Ardenne",
    "franche comte": "Franche-Comté",
    "bourgogne franche comte": "Bourgogne-Franche-Comté",
    "auvergne rhone alpes": "Auvergne-Rhône-Alpes",
    "rhone alpes": "Rhône-Alpes",
    "basse normandie": "Basse-Normandie",
    "haute normandie": "Haute-Normandie",
    "hauts de france": "Hauts-de-France",
    "nouvelle aquitaine": "Nouvelle-Aquitaine",
    "grand est": "Grand Est",
}

_PERIODES = ("Nicheurs", "Hivernants", "De passage")
_ANNEE = re.compile(r"\b(19|20)\d{2}\b")


def _net(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").replace("\n", " ")).strip()


def _cle_territoire(libelle: str) -> str:
    """Forme comparable d'un nom de territoire (accents et tirets neutralisés)."""
    import unicodedata
    s = unicodedata.normalize("NFD", str(libelle or "")).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9 ]", " ", s.lower().replace("-", " "))
    return re.sub(r"\s+", " ", s).strip()


def territoire_bdc(libelle: str) -> str:
    """Nom du territoire tel que BDC-Statuts l'écrit."""
    return ALIAS_TERRITOIRE.get(_cle_territoire(libelle), _net(libelle))


@dataclass
class ColonneStatut:
    index: int               # index 0-based dans l'onglet Statuts
    champ: str               # LRE, PN, LRF Nicheurs…
    type_bdc: str            # LRE, PN, LRN, LRR, ZDET…
    territoire: str          # tel qu'écrit par Marie
    territoire_bdc: str      # tel qu'écrit par BDC-Statuts
    periode: str = ""        # Nicheurs | Hivernants | De passage | ""
    annee: str = ""          # édition, quand l'en-tête la porte
    libelle: str = ""        # en-tête brut dans Statuts

    @property
    def comparable(self) -> bool:
        """Colonne réellement confrontée à BDC-Statuts.

        Trois conditions, et il faut les trois :

        * la colonne porte un type reconnu ;
        * elle n'est pas périodisée — BDC ne distingue pas nicheurs, hivernants
          et de passage, confronter « nicheurs » à un statut global n'aurait
          pas de sens ;
        * les deux vocabulaires se rejoignent. `PN`, `PR` et `DO` échouent ici :
          BDC désigne l'arrêté (`NO3`, `RV93`) ou l'annexe (`CDO1`) là où le
          classeur écrit l'article (« Art. 3 ») ou le code espèce (`A092`).

        La troisième a été oubliée un temps : le décompte affiché annonçait 11
        colonnes pour les oiseaux alors que 9 seulement étaient confrontées.
        """
        from . import vocabulaire
        return (bool(self.type_bdc) and not self.periode
                and vocabulaire.comparable(self.type_bdc))


@dataclass
class LigneTaxon:
    ligne: int               # numéro de ligne dans l'onglet Statuts
    cd_nom: str
    nom_scientifique: str
    nom_commun: str
    auteur: str = ""
    valeurs: dict[int, str] = field(default_factory=dict)   # index colonne -> valeur

    @property
    def libelle(self) -> str:
        return self.nom_commun or self.nom_scientifique or f"(ligne {self.ligne})"


@dataclass
class Classeur:
    groupe: str
    chemin: Path
    taxons: list[LigneTaxon]
    colonnes: list[ColonneStatut]
    avertissements: list[str] = field(default_factory=list)

    def colonne(self, type_bdc: str, territoire_bdc: str | None = None) -> ColonneStatut | None:
        for c in self.colonnes:
            if c.type_bdc == type_bdc and not c.periode:
                if territoire_bdc is None or c.territoire_bdc == territoire_bdc:
                    return c
        return None


def _dictionnaire_depuis_saisie(wb) -> tuple[list[ColonneStatut], list[str]]:
    """Lit les lignes 5/6/7 de « Saisie étude » : index, territoire, champ."""
    nom = next((n for n in wb.sheetnames if n.strip().lower().startswith("saisie")), None)
    if nom is None:
        return [], ["onglet « Saisie étude » absent"]

    ws = wb[nom]
    lignes = {i: [_net(c.value) for c in ws[i]] for i in (5, 6, 7)}
    if not any(v.isdigit() for v in lignes[5]):
        return [], ["« Saisie étude » ne porte pas les index de colonnes en ligne 5"]

    colonnes: list[ColonneStatut] = []
    territoire_courant = ""
    for i, brut in enumerate(lignes[5]):
        if not brut.isdigit():
            continue
        champ = lignes[7][i] if i < len(lignes[7]) else ""
        if not champ:
            continue
        terr = lignes[6][i] if i < len(lignes[6]) else ""
        if terr:
            territoire_courant = terr
        terr = terr or territoire_courant

        periode = next((p for p in _PERIODES if p.lower() in champ.lower()), "")
        racine = re.sub(r"(?i)\s*(nicheurs|hivernants|de passage)\s*", " ", champ).strip()
        type_bdc = CHAMP_VERS_BDC.get(racine.upper(), "")
        if not type_bdc and racine.upper().startswith("LRR"):
            type_bdc = "LRR"

        colonnes.append(ColonneStatut(
            index=int(brut) - 1, champ=champ, type_bdc=type_bdc,
            territoire=terr, territoire_bdc=territoire_bdc(terr),
            periode=periode,
        ))
    return colonnes, []


def _dictionnaire_depuis_statuts(ws) -> tuple[list[ColonneStatut], list[str]]:
    """Repli : en-têtes de `Statuts`, plages fusionnées propagées."""
    entetes = [_net(c.value) for c in ws[1]]
    fusions: dict[int, tuple[str, int, int]] = {}
    for plage in ws.merged_cells.ranges:
        if plage.min_row != 1:
            continue
        libelle = _net(ws.cell(1, plage.min_col).value)
        largeur = plage.max_col - plage.min_col + 1
        for rang, col in enumerate(range(plage.min_col, plage.max_col + 1)):
            fusions[col - 1] = (libelle, rang, largeur)

    colonnes: list[ColonneStatut] = []
    avert: list[str] = []
    for i, brut in enumerate(entetes):
        libelle, rang, largeur = fusions.get(i, (brut, 0, 1))
        if not libelle:
            continue
        bas = libelle.lower()
        if bas.startswith("liste rouge europe"):
            type_bdc, terr = "LRE", "Europe"
        elif bas.startswith("liste rouge france"):
            type_bdc, terr = "LRN", "France"
        elif bas.startswith("liste rouge"):
            type_bdc, terr = "LRR", _ANNEE.sub("", libelle[len("liste rouge"):]).strip()
        elif bas.startswith("dz "):
            type_bdc, terr = "ZDET", libelle[3:].strip()
        elif bas.startswith("protection nationale"):
            type_bdc, terr = "PN", "France"
        elif bas.startswith("protection "):
            type_bdc, terr = "PR", libelle[len("protection "):].strip()
        elif "annexe i " in bas and "oiseaux" in bas:
            type_bdc, terr = "DO", "Europe"
        elif "annexe ii" in bas:
            type_bdc, terr = "DH", "Europe"
        else:
            continue

        periode = ""
        if largeur > 1:
            periode = _PERIODES[rang] if rang < len(_PERIODES) else f"colonne {rang + 1}"
            if rang == 0:
                avert.append(
                    f"« {libelle} » couvre {largeur} colonnes : périodes déduites de "
                    "leur ordre, à confirmer"
                )
        annee = (_ANNEE.search(libelle).group(0) if _ANNEE.search(libelle) else "")
        colonnes.append(ColonneStatut(
            index=i, champ=libelle, type_bdc=type_bdc, territoire=terr,
            territoire_bdc=territoire_bdc(terr), periode=periode, annee=annee,
            libelle=libelle,
        ))
    return colonnes, avert


def lire(chemin: Path, groupe: str) -> Classeur:
    """Charge un classeur B-Statuts. Aucune écriture, jamais."""
    wb = openpyxl.load_workbook(chemin, data_only=True)
    if "Statuts" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{chemin.name} : onglet « Statuts » introuvable")
    ws = wb["Statuts"]

    colonnes, avert = _dictionnaire_depuis_saisie(wb)
    if not colonnes:
        colonnes, avert2 = _dictionnaire_depuis_statuts(ws)
        avert = avert + avert2

    # Les en-têtes de Statuts complètent le dictionnaire (libellé brut, année
    # d'édition quand Marie l'a notée : « Liste rouge Pays de la Loire 2014 »).
    entetes = [_net(c.value) for c in ws[1]]
    for c in colonnes:
        if 0 <= c.index < len(entetes):
            c.libelle = c.libelle or entetes[c.index]
            trouve = _ANNEE.search(entetes[c.index])
            c.annee = c.annee or (trouve.group(0) if trouve else "")

    hauts = [h.upper() for h in entetes]
    try:
        i_cd = hauts.index("CD_NOM")
    except ValueError:
        wb.close()
        raise ValueError(f"{chemin.name} : colonne CD_NOM introuvable") from None
    i_sci = next((i for i, h in enumerate(hauts) if h.startswith("NOM SCIENTIFIQUE")), None)
    i_com = next((i for i, h in enumerate(hauts) if h.startswith("NOM COMMUN")), None)
    i_aut = next((i for i, h in enumerate(hauts) if h.startswith("AUTEUR")), None)

    utiles = sorted({c.index for c in colonnes})
    taxons: list[LigneTaxon] = []
    for n, brut in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def at(i):
            return _net(brut[i]) if i is not None and i < len(brut) else ""
        cd, sci, com = at(i_cd), at(i_sci), at(i_com)
        if not cd and not sci and not com:
            continue
        taxons.append(LigneTaxon(
            ligne=n, cd_nom=cd, nom_scientifique=sci, nom_commun=com, auteur=at(i_aut),
            valeurs={i: at(i) for i in utiles if at(i)},
        ))
    wb.close()
    return Classeur(groupe=groupe, chemin=chemin, taxons=taxons,
                    colonnes=colonnes, avertissements=avert)


def lire_tous(dossier: Path, correspondances: dict[str, str]) -> dict[str, Classeur]:
    """Charge les classeurs présents dans `dossier` selon `{groupe: motif}`."""
    sortie: dict[str, Classeur] = {}
    for groupe, motif in correspondances.items():
        trouves = sorted(dossier.glob(motif))
        if not trouves:
            continue
        if len(trouves) > 1:
            trouves = [max(trouves, key=lambda p: p.stat().st_mtime)]
        sortie[groupe] = lire(trouves[0], groupe)
    return sortie
